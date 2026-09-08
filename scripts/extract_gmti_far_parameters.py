#!/usr/bin/env python3
"""从指定工作簿提取“广域GMTI（远距离）”参数并校验派生量。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


C = 299_792_458.0


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    raw = args.input.read_bytes()
    sha256 = hashlib.sha256(raw).hexdigest()
    workbook = load_workbook(args.input, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    columns = {
        str(sheet.cell(2, column).value).strip(): column
        for column in range(2, sheet.max_column + 1)
        if sheet.cell(2, column).value is not None
    }
    mode = "广域GMTI（远距离）"
    if mode not in columns:
        raise SystemExit(f"工作簿缺少{mode!r}列")
    column = columns[mode]
    values = {str(sheet.cell(row, 1).value).strip(): sheet.cell(row, column).value
              for row in range(3, sheet.max_row + 1)
              if sheet.cell(row, 1).value is not None}
    required = ["分辨率", "带宽(Mhz)", "幅宽(m)", "幅宽时间(us)", "脉宽(us)",
                "接收总时长(us)", "DDC后采样率(Mhz)", "采样点数", "FPGA上传点数",
                "接收延时", "prt", "prf", "占空比", "数据传输速率GB/s", "对应扫描速度°/s"]
    missing = [key for key in required if values.get(key) is None]
    if missing:
        raise SystemExit("工作簿GMTI远距离列缺值：" + ", ".join(missing))

    fs_mhz = float(values["DDC后采样率(Mhz)"])
    pulse_us = float(values["脉宽(us)"])
    acquired = int(values["采样点数"])
    uploaded = int(values["FPGA上传点数"])
    delay_us = float(values["接收延时"])
    pulse_samples = round(pulse_us * fs_mhz)
    valid_samples = acquired - pulse_samples
    if acquired != round(float(values["接收总时长(us)"]) * fs_mhz):
        raise SystemExit("采样点数与接收总时长×采样率不一致")
    if uploaded < acquired or uploaded % 64:
        raise SystemExit("FPGA上传点数必须不小于实际采样点数且为64的倍数")
    if not math.isclose(valid_samples / fs_mhz, float(values["幅宽时间(us)"]), abs_tol=1e-9):
        raise SystemExit("幅宽时间与（采样点数-脉宽点数）/采样率不一致")

    nearest = 0.5 * C * delay_us * 1e-6
    swath = 0.5 * C * valid_samples / (fs_mhz * 1e6)
    payload = {
        "source": {"path": str(args.input.resolve()), "sha256": sha256,
                   "sheet": sheet.title, "column": sheet.cell(2, column).column_letter,
                   "mode": mode, "extracted_on": date.today().isoformat()},
        "workbook_values": {
            "range_resolution_m": float(str(values["分辨率"]).split("m")[0]),
            "bandwidth_mhz": float(values["带宽(Mhz)"]),
            "swath_m": float(values["幅宽(m)"]),
            "swath_time_us": float(values["幅宽时间(us)"]),
            "pulse_width_us": pulse_us,
            "receive_duration_us": float(values["接收总时长(us)"]),
            "ddc_sample_rate_mhz": fs_mhz, "acquired_samples": acquired,
            "fpga_uploaded_samples": uploaded, "receive_delay_us": delay_us,
            "prt_us": float(values["prt"]), "prf_hz": float(values["prf"]),
            "duty_cycle": float(values["占空比"]),
            "transfer_rate_gib_per_s": float(values["数据传输速率GB/s"]),
            "scan_speed_deg_per_s": float(str(values["对应扫描速度°/s"]).replace("°/s", "")),
            "scan_ranges_deg": [45.0, 90.0, 180.0]},
        "derived_processing": {
            "pulse_width_samples": pulse_samples, "valid_swath_samples": valid_samples,
            "fpga_padding_samples": uploaded - acquired, "range_crop_start": 0,
            "range_compress_len": valid_samples, "range_fft_len": 12288,
            "nearest_slant_range_m": nearest, "valid_swath_m": swath,
            "farthest_slant_range_m": nearest + swath,
            "derivation": "c=299792458m/s; Rnear=c*receive_delay/2; swath=c*(acquired_samples/fs-pulse_width)/2"},
        "non_workbook_inputs": {"pulse_count_per_beam": 128,
            "pulse_count_source": "新协议现有生产接口；工作簿N20为空",
            "beam_count": 61, "scan_min_deg": -60.0, "scan_step_deg": 2.0}}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
                           encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
